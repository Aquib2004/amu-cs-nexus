"""Private, expiring file Q&A for YouRobo.

The original binary is processed in memory and discarded. Only extracted text,
chunks, metadata, and a hash of a random access token are persisted.
"""

import hashlib
import hmac
import io
import re
import secrets
from datetime import datetime, timedelta, timezone
from pathlib import Path
from uuid import UUID

from sqlalchemy import delete
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.chat_upload import ChatUpload, ChatUploadChunk
from app.services.knowledge import KnowledgeHit, _score, _tokens

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_TEXT_CHARS = 1_000_000
MAX_CHUNKS = 1_000
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}
ALLOWED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/plain",
    "text/markdown",
}


class UploadValidationError(ValueError):
    pass


class UploadAccessError(PermissionError):
    pass


def _token_hash(token: str) -> str:
    return hashlib.sha256(token.encode("utf-8")).hexdigest()


def _safe_filename(filename: str) -> str:
    name = Path(filename or "upload").name
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip()
    return (name or "upload")[:255]


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise UploadValidationError("The PDF is password-protected.") from exc
    pages: list[str] = []
    for index, page in enumerate(reader.pages[:400], start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[Page {index}]\n{text}")
        if sum(len(part) for part in pages) >= MAX_TEXT_CHARS:
            break
    return "\n\n".join(pages)[:MAX_TEXT_CHARS]


def _extract_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            value = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if value:
                parts.append(value)
    return "\n".join(parts)[:MAX_TEXT_CHARS]


def _chunk_text(text: str, size: int = 1200) -> list[str]:
    paragraphs = re.split(r"\n{2,}|(?<=[.!?])\s+", text)
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if len(paragraph) > size:
            if current:
                chunks.append(current)
                current = ""
            chunks.extend(paragraph[i:i + size] for i in range(0, len(paragraph), size))
        elif len(current) + len(paragraph) + 1 <= size:
            current = f"{current} {paragraph}".strip()
        else:
            if current:
                chunks.append(current)
            current = paragraph
        if len(chunks) >= MAX_CHUNKS:
            break
    if current and len(chunks) < MAX_CHUNKS:
        chunks.append(current)
    return chunks[:MAX_CHUNKS]

def extract_text(filename: str, content_type: str, data: bytes) -> str:
    name = _safe_filename(filename)
    extension = Path(name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise UploadValidationError("Supported files: PDF, DOCX, TXT, and Markdown.")
    if not data:
        raise UploadValidationError("The uploaded file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise UploadValidationError("The maximum file size is 10 MB.")
    normalized_type = (content_type or "").split(";", 1)[0].lower()
    if normalized_type and normalized_type not in ALLOWED_CONTENT_TYPES and not normalized_type.startswith("text/"):
        raise UploadValidationError("The uploaded file type is not supported.")
    try:
        if extension == ".pdf":
            text = _extract_pdf(data)
        elif extension == ".docx":
            text = _extract_docx(data)
        else:
            text = data.decode("utf-8-sig", errors="replace")
    except UploadValidationError:
        raise
    except Exception as exc:
        raise UploadValidationError("The file could not be read or contains no extractable text.") from exc
    text = text.replace("\x00", " ").strip()
    if len(text) < 20:
        raise UploadValidationError("No useful text could be extracted from this file.")
    return text[:MAX_TEXT_CHARS]

def _embed(texts: list[str]) -> list[list[float]] | None:
    key = settings.embedding_api_key or settings.llm_api_key or settings.gemini_api_key
    if not key:
        return None
    try:
        from ai.providers.gemini_embed import GeminiEmbedder

        embedder = GeminiEmbedder(
            api_key=key,
            timeout=settings.llm_timeout_seconds,
            max_retries=settings.llm_max_retries,
        )
        vectors: list[list[float]] = []
        for offset in range(0, len(texts), 64):
            vectors.extend(embedder.embed_many(texts[offset:offset + 64]))
        return vectors
    except Exception:
        return None


def create_upload(session: Session, filename: str, content_type: str, data: bytes) -> tuple[ChatUpload, str]:
    text = extract_text(filename, content_type, data)
    token = secrets.token_urlsafe(32)
    upload = ChatUpload(
        access_token_hash=_token_hash(token),
        filename=_safe_filename(filename),
        content_type=(content_type or "application/octet-stream")[:150],
        size_bytes=len(data),
        status="ready",
        expires_at=datetime.now(timezone.utc) + timedelta(hours=settings.upload_expiry_hours),
    )
    pieces = _chunk_text(text)
    vectors = _embed(pieces)
    for index, piece in enumerate(pieces):
        upload.chunks.append(ChatUploadChunk(
            chunk_index=index,
            text=piece,
            embedding=vectors[index] if vectors else None,
        ))
    session.add(upload)
    session.commit()
    return upload, token


def get_authorized_upload(session: Session, upload_id: UUID, token: str | None) -> ChatUpload:
    upload = session.get(ChatUpload, upload_id)
    expected = _token_hash(token or "")
    if upload is None or not hmac.compare_digest(upload.access_token_hash, expected):
        raise UploadAccessError("Upload not found or access token is invalid.")
    expires = upload.expires_at
    if expires.tzinfo is None:
        expires = expires.replace(tzinfo=timezone.utc)
    if expires <= datetime.now(timezone.utc):
        session.delete(upload)
        session.commit()
        raise UploadAccessError("This upload has expired. Please upload the file again.")
    return upload


def search_upload(upload: ChatUpload, query: str, limit: int = 8) -> list[KnowledgeHit]:
    query_vector = None
    key = settings.embedding_api_key or settings.llm_api_key or settings.gemini_api_key
    if key and any(chunk.embedding for chunk in upload.chunks):
        try:
            from ai.providers.gemini_embed import GeminiEmbedder
            query_vector = GeminiEmbedder(api_key=key).embed(query)
        except Exception:
            query_vector = None
    scored: list[tuple[ChatUploadChunk, float]] = []
    for chunk in upload.chunks:
        score = _score(query, chunk.text)
        if query_vector and chunk.embedding and len(chunk.embedding) == len(query_vector):
            from ai.retrieval.scoring import cosine_similarity
            score += max(0.0, cosine_similarity(query_vector, chunk.embedding))
        if score > 0:
            scored.append((chunk, score))
    scored.sort(key=lambda pair: pair[1], reverse=True)
    return [
        KnowledgeHit(
            chunk_id=chunk.id,
            document_id=upload.id,
            text=chunk.text,
            source_url=f"upload://{upload.filename}",
            score=score,
            title=upload.filename,
            source_type="upload",
        )
        for chunk, score in scored[:limit]
    ]


def delete_upload(session: Session, upload_id: UUID, token: str) -> None:
    upload = get_authorized_upload(session, upload_id, token)
    session.delete(upload)
    session.commit()


def purge_expired_uploads(session: Session) -> int:
    result = session.execute(delete(ChatUpload).where(ChatUpload.expires_at <= datetime.now(timezone.utc)))
    session.commit()
    return int(result.rowcount or 0)
